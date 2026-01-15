"""Excel parsing utility for xlsx and xls files (read-only)"""

import io
from pathlib import Path
from typing import Optional, Any, Iterator
from dataclasses import dataclass, field

import openpyxl
import xlrd
import msoffcrypto


@dataclass
class ExcelSheet:
    """single sheet data"""

    name: str
    rows: list[list[Any]] = field(default_factory=list)
    row_count: int = 0
    col_count: int = 0

    def to_dicts(self, header_row: int = 0) -> list[dict[str, Any]]:
        """convert rows to list of dicts using specified row as header

        Args:
            header_row: row index to use as header (0-based)

        Returns:
            list of dicts with header values as keys
        """
        if not self.rows or header_row >= len(self.rows):
            return []

        headers = [
            str(h) if h is not None else f"col_{i}"
            for i, h in enumerate(self.rows[header_row])
        ]

        result = []
        for row in self.rows[header_row + 1 :]:
            row_dict = {}
            for i, cell in enumerate(row):
                key = headers[i] if i < len(headers) else f"col_{i}"
                row_dict[key] = cell
            result.append(row_dict)
        return result


@dataclass
class ExcelDocument:
    """parsed Excel document"""

    file_path: str
    sheet_names: list[str] = field(default_factory=list)
    sheets: dict[str, ExcelSheet] = field(default_factory=dict)

    def get_sheet(self, name: str) -> Optional[ExcelSheet]:
        """get sheet by name"""
        return self.sheets.get(name)

    def get_sheet_by_index(self, index: int) -> Optional[ExcelSheet]:
        """get sheet by index"""
        if 0 <= index < len(self.sheet_names):
            name = self.sheet_names[index]
            return self.sheets.get(name)
        return None

    @property
    def first_sheet(self) -> Optional[ExcelSheet]:
        """get first sheet"""
        return self.get_sheet_by_index(0)


class ExcelParser:
    """Excel parser supporting both xlsx and xls formats (read-only)

    Usage:
        # parse entire workbook
        parser = ExcelParser("path/to/file.xlsx")
        doc = parser.parse()

        # parse with password
        parser = ExcelParser("encrypted.xlsx", password="secret")
        doc = parser.parse()

        # access sheets
        sheet = doc.first_sheet
        print(sheet.rows)

        # convert to dicts
        data = sheet.to_dicts()

        # parse specific sheets
        doc = parser.parse(sheets=["Sheet1", "Sheet2"])

        # parse from bytes
        parser = ExcelParser.from_bytes(excel_bytes, filename="data.xlsx", password="secret")
        doc = parser.parse()
    """

    def __init__(self, file_path: str | Path, password: Optional[str] = None):
        """initialize parser with file path

        Args:
            file_path: path to Excel file (.xlsx or .xls)
            password: password for encrypted file (optional)
        """
        self.file_path = Path(file_path)
        self._password = password
        self._is_xls = self.file_path.suffix.lower() == ".xls"
        self._workbook: Any = None
        self._bytes_io: Optional[io.BytesIO] = None
        self._decrypted_io: Optional[io.BytesIO] = None

    @classmethod
    def from_bytes(
        cls,
        excel_bytes: bytes,
        filename: str = "memory.xlsx",
        password: Optional[str] = None,
    ) -> "ExcelParser":
        """create parser from bytes (e.g., uploaded file)

        Args:
            excel_bytes: Excel file content as bytes
            filename: virtual filename (used to detect format by extension)
            password: password for encrypted file (optional)

        Returns:
            ExcelParser instance
        """
        parser = cls.__new__(cls)
        parser.file_path = Path(filename)
        parser._password = password
        parser._is_xls = parser.file_path.suffix.lower() == ".xls"
        parser._bytes_io = io.BytesIO(excel_bytes)
        parser._decrypted_io = None
        parser._workbook = None
        return parser

    def _decrypt_file(self, file_obj: io.BytesIO) -> io.BytesIO:
        """decrypt encrypted Excel file using msoffcrypto

        Args:
            file_obj: encrypted file as BytesIO

        Returns:
            decrypted file as BytesIO
        """
        file_obj.seek(0)
        office_file = msoffcrypto.OfficeFile(file_obj)

        if not office_file.is_encrypted():
            file_obj.seek(0)
            return file_obj

        decrypted = io.BytesIO()
        office_file.load_key(password=self._password or "")
        office_file.decrypt(decrypted)
        decrypted.seek(0)
        return decrypted

    def _open(self):
        """open Excel file (read-only)"""
        if self._workbook is not None:
            return self._workbook

        # prepare file source
        if self._bytes_io:
            file_source = self._bytes_io
        else:
            file_source = io.BytesIO(self.file_path.read_bytes())

        # decrypt if password provided
        if self._password:
            self._decrypted_io = self._decrypt_file(file_source)
            file_source = self._decrypted_io

        if self._is_xls:
            # xls format - use xlrd (read-only by nature)
            file_source.seek(0)
            self._workbook = xlrd.open_workbook(file_contents=file_source.getvalue())
        else:
            # xlsx format - use openpyxl in read-only mode
            file_source.seek(0)
            self._workbook = openpyxl.load_workbook(
                file_source, data_only=True, read_only=True
            )

        return self._workbook

    def close(self):
        """close workbook"""
        if self._workbook:
            if not self._is_xls:
                self._workbook.close()
            self._workbook = None
        if self._bytes_io:
            self._bytes_io.close()
            self._bytes_io = None
        if self._decrypted_io:
            self._decrypted_io.close()
            self._decrypted_io = None

    def __enter__(self):
        self._open()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()

    def _get_sheet_names(self) -> list[str]:
        """get all sheet names"""
        wb = self._open()
        if self._is_xls:
            return wb.sheet_names()
        else:
            return wb.sheetnames

    def _parse_xls_sheet(self, sheet) -> ExcelSheet:
        """parse xls sheet using xlrd"""
        rows = []
        for row_idx in range(sheet.nrows):
            row = []
            for col_idx in range(sheet.ncols):
                cell = sheet.cell(row_idx, col_idx)
                value = cell.value

                # handle date cells
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, self._workbook.datemode)
                    except Exception:
                        pass

                row.append(value)
            rows.append(row)

        return ExcelSheet(
            name=sheet.name,
            rows=rows,
            row_count=sheet.nrows,
            col_count=sheet.ncols,
        )

    def _parse_xlsx_sheet(self, sheet) -> ExcelSheet:
        """parse xlsx sheet using openpyxl"""
        rows = []
        max_col = 0

        for row in sheet.iter_rows():
            row_data = [cell.value for cell in row]
            rows.append(row_data)
            max_col = max(max_col, len(row_data))

        return ExcelSheet(
            name=sheet.title,
            rows=rows,
            row_count=len(rows),
            col_count=max_col,
        )

    def parse(
        self,
        sheets: Optional[list[str | int]] = None,
        skip_empty_rows: bool = False,
    ) -> ExcelDocument:
        """parse Excel document

        Args:
            sheets: list of sheet names or indices to parse, None for all
            skip_empty_rows: whether to skip completely empty rows

        Returns:
            ExcelDocument with parsed content
        """
        wb = self._open()
        sheet_names = self._get_sheet_names()

        doc = ExcelDocument(
            file_path=str(self.file_path),
            sheet_names=sheet_names,
        )

        # determine which sheets to parse
        sheets_to_parse: list[str] = []
        if sheets is None:
            sheets_to_parse = sheet_names
        else:
            for s in sheets:
                if isinstance(s, int):
                    if 0 <= s < len(sheet_names):
                        sheets_to_parse.append(sheet_names[s])
                elif s in sheet_names:
                    sheets_to_parse.append(s)

        # parse each sheet
        for sheet_name in sheets_to_parse:
            if self._is_xls:
                sheet = wb.sheet_by_name(sheet_name)
                excel_sheet = self._parse_xls_sheet(sheet)
            else:
                sheet = wb[sheet_name]
                excel_sheet = self._parse_xlsx_sheet(sheet)

            # optionally skip empty rows
            if skip_empty_rows:
                excel_sheet.rows = [
                    row
                    for row in excel_sheet.rows
                    if any(cell is not None and cell != "" for cell in row)
                ]
                excel_sheet.row_count = len(excel_sheet.rows)

            doc.sheets[sheet_name] = excel_sheet

        return doc

    def get_sheet_names(self) -> list[str]:
        """get all sheet names without full parsing"""
        self._open()
        return self._get_sheet_names()

    def iter_rows(
        self,
        sheet: str | int = 0,
        start_row: int = 0,
        end_row: Optional[int] = None,
    ) -> Iterator[list[Any]]:
        """iterate rows lazily (memory efficient for large files)

        Args:
            sheet: sheet name or index
            start_row: starting row (0-based)
            end_row: ending row (exclusive), None for all

        Yields:
            row data as list
        """
        wb = self._open()
        sheet_names = self._get_sheet_names()

        # resolve sheet name
        if isinstance(sheet, int):
            if 0 <= sheet < len(sheet_names):
                sheet_name = sheet_names[sheet]
            else:
                return
        else:
            sheet_name = sheet

        if self._is_xls:
            ws = wb.sheet_by_name(sheet_name)
            end = end_row if end_row else ws.nrows
            for row_idx in range(start_row, min(end, ws.nrows)):
                row = []
                for col_idx in range(ws.ncols):
                    cell = ws.cell(row_idx, col_idx)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate_as_datetime(value, wb.datemode)
                        except Exception:
                            pass
                    row.append(value)
                yield row
        else:
            ws = wb[sheet_name]
            for row_idx, row in enumerate(ws.iter_rows()):
                if row_idx < start_row:
                    continue
                if end_row and row_idx >= end_row:
                    break
                yield [cell.value for cell in row]


# convenience functions
def parse_excel(
    file_path: str | Path,
    password: Optional[str] = None,
    **kwargs,
) -> ExcelDocument:
    """quick parse an Excel file

    Args:
        file_path: path to Excel file
        password: password for encrypted file (optional)
        **kwargs: additional arguments passed to parse()

    Returns:
        ExcelDocument with parsed content
    """
    with ExcelParser(file_path, password=password) as parser:
        return parser.parse(**kwargs)


def parse_excel_bytes(
    excel_bytes: bytes,
    filename: str = "data.xlsx",
    password: Optional[str] = None,
    **kwargs,
) -> ExcelDocument:
    """quick parse Excel from bytes

    Args:
        excel_bytes: Excel content as bytes
        filename: virtual filename (determines format by extension)
        password: password for encrypted file (optional)
        **kwargs: additional arguments passed to parse()

    Returns:
        ExcelDocument with parsed content
    """
    parser = ExcelParser.from_bytes(excel_bytes, filename, password=password)
    try:
        return parser.parse(**kwargs)
    finally:
        parser.close()


def excel_to_dicts(
    file_path: str | Path,
    sheet: str | int = 0,
    header_row: int = 0,
    password: Optional[str] = None,
) -> list[dict[str, Any]]:
    """quick convert Excel sheet to list of dicts

    Args:
        file_path: path to Excel file
        sheet: sheet name or index
        header_row: row index to use as header
        password: password for encrypted file (optional)

    Returns:
        list of dicts
    """
    with ExcelParser(file_path, password=password) as parser:
        if isinstance(sheet, int):
            doc = parser.parse(sheets=[sheet])
            excel_sheet = doc.get_sheet_by_index(0)
        else:
            doc = parser.parse(sheets=[sheet])
            excel_sheet = doc.get_sheet(sheet)

        if excel_sheet:
            return excel_sheet.to_dicts(header_row)
        return []
